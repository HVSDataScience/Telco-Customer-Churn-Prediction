"""
Generate project_report.docx for Telco Customer Churn Prediction project.
Run: python generate_report.py
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime


def set_heading_style(paragraph, level=1):
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1f, 0x23, 0x28)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x3b, 0x82, 0xd4)
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x57, 0x60, 0x6a)


def add_section_title(doc, title, level=1):
    para = doc.add_heading(title, level=level)
    para.runs[0].font.color.rgb = (
        RGBColor(0x1f, 0x23, 0x28) if level == 1 else RGBColor(0x3b, 0x82, 0xd4)
    )
    para.runs[0].bold = True
    return para


def add_body(doc, text):
    para = doc.add_paragraph(text)
    para.runs[0].font.size = Pt(11) if para.runs else None
    return para


def add_bullet(doc, text):
    para = doc.add_paragraph(text, style='List Bullet')
    return para


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'

    # Header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        run = hdr_cells[i].paragraphs[0].runs[0]
        run.bold = True
        run.font.size = Pt(10)
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Data rows
    for r_idx, row in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, cell_val in enumerate(row):
            row_cells[c_idx].text = str(cell_val)
            row_cells[c_idx].paragraphs[0].runs[0].font.size = Pt(10)
            row_cells[c_idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    return table


# ─────────────────────────────────────────────────────────────────────────────
doc = Document()

# Set default font
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# ── Title Page ────────────────────────────────────────────────────────────────
doc.add_paragraph()
title_para = doc.add_paragraph()
title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title_para.add_run('Telco Customer Churn Prediction')
title_run.bold = True
title_run.font.size = Pt(22)
title_run.font.color.rgb = RGBColor(0x1f, 0x23, 0x28)

sub_para = doc.add_paragraph()
sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub_para.add_run('Machine Learning Project Report')
sub_run.font.size = Pt(14)
sub_run.font.color.rgb = RGBColor(0x3b, 0x82, 0xd4)

doc.add_paragraph()
date_para = doc.add_paragraph()
date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
date_para.add_run(f'Date: {datetime.date.today().strftime("%B %d, %Y")}').font.size = Pt(11)

doc.add_page_break()

# ── Abstract ──────────────────────────────────────────────────────────────────
add_section_title(doc, 'Abstract', level=1)
add_body(doc,
    'This report presents a machine learning project for predicting customer churn in the '
    'telecommunications industry using the IBM Telco Customer Churn dataset. The dataset '
    'contains 7,032 cleaned records with 20 features including contract type, internet '
    'service, payment method, tenure, and monthly charges. Two classification models — '
    'Logistic Regression and Random Forest — were trained and evaluated. Logistic Regression '
    'achieved the best performance with an Accuracy of 79.39%, Recall of 56.42%, and '
    'ROC-AUC of 83.45%, and was selected as the final model. Key findings show that '
    'month-to-month contracts (42.71% churn), fiber optic internet (41.89%), electronic '
    'check payment (45.29%), and short tenure (47.68% churn for customers under 12 months) '
    'are the strongest churn indicators. Actionable business recommendations are provided '
    'to support proactive customer retention strategies.'
)
doc.add_paragraph()

# ── 1. Introduction ───────────────────────────────────────────────────────────
add_section_title(doc, '1. Introduction', level=1)
add_body(doc,
    'Customer churn — the rate at which customers stop doing business with a company — '
    'is one of the most critical challenges in the telecommunications industry. '
    'Acquiring a new customer is significantly more expensive than retaining an existing one, '
    'making churn prediction a high-value business problem. This project uses the IBM Telco '
    'Customer Churn dataset to build, evaluate, and compare machine learning models that '
    'identify customers at risk of churning.'
)
doc.add_paragraph()

# ── 2. Problem Statement ──────────────────────────────────────────────────────
add_section_title(doc, '2. Problem Statement', level=1)
add_body(doc,
    'A telecom company faces a churn rate of 26.58%, meaning more than one in four customers '
    'leaves the company. Without an early-warning system, the company cannot proactively '
    'intervene. The problem is to build a binary classification model that accurately predicts '
    'whether a customer will churn (Yes/No) based on their demographics, account information, '
    'and services subscribed.'
)
doc.add_paragraph()

# ── 3. Objective ──────────────────────────────────────────────────────────────
add_section_title(doc, '3. Objective', level=1)
bullets = [
    'Load, clean, and explore the Telco Customer Churn dataset.',
    'Analyze churn patterns across key business dimensions.',
    'Build and evaluate Logistic Regression and Random Forest models.',
    'Select the best-performing model based on real evaluation metrics.',
    'Generate customer churn predictions and churn probabilities.',
    'Provide actionable business recommendations to reduce churn.',
]
for b in bullets:
    add_bullet(doc, b)
doc.add_paragraph()

# ── 4. Dataset ────────────────────────────────────────────────────────────────
add_section_title(doc, '4. Dataset', level=1)
add_body(doc,
    'Dataset: IBM Telco Customer Churn\n'
    'Source: Kaggle (https://www.kaggle.com/datasets/blastchar/telco-customer-churn)\n'
    'Original records: 7,043 | After cleaning: 7,032 | Features: 20 + 1 target (Churn)'
)
doc.add_paragraph()

add_section_title(doc, '4.1 Key Features', level=2)
feat_headers = ['Feature', 'Type', 'Description']
feat_rows = [
    ('customerID', 'Identifier', 'Unique customer identifier (dropped)'),
    ('gender', 'Categorical', 'Customer gender (Male/Female)'),
    ('SeniorCitizen', 'Binary', 'Whether the customer is a senior citizen'),
    ('tenure', 'Numerical', 'Number of months with the company'),
    ('Contract', 'Categorical', 'Contract type (Month-to-month, One year, Two year)'),
    ('MonthlyCharges', 'Numerical', 'Monthly amount charged to the customer'),
    ('TotalCharges', 'Numerical', 'Total amount charged over customer lifetime'),
    ('InternetService', 'Categorical', 'Internet service type (DSL, Fiber optic, No)'),
    ('PaymentMethod', 'Categorical', 'Payment method used'),
    ('Churn', 'Target', 'Whether the customer churned (Yes/No)'),
]
add_table(doc, feat_headers, feat_rows)

# ── 5. Methodology ────────────────────────────────────────────────────────────
add_section_title(doc, '5. Methodology', level=1)
add_body(doc,
    'The project follows a standard machine learning pipeline: data loading, '
    'cleaning, exploratory analysis, feature engineering, model training, '
    'evaluation, and selection. Two classification algorithms were trained and '
    'compared — Logistic Regression and Random Forest. The dataset was split '
    '80% for training and 20% for testing using stratified sampling to preserve '
    'the class distribution.'
)
doc.add_paragraph()

# ── 6. Data Cleaning ──────────────────────────────────────────────────────────
add_section_title(doc, '6. Data Cleaning', level=1)
add_body(doc, 'The following steps were applied to clean the dataset:')
cleaning_steps = [
    'TotalCharges column: Contained 11 blank string entries (new customers with 0 tenure). '
     'Converted to numeric using pd.to_numeric(errors="coerce"), resulting in 11 NaN values.',
    'Dropped 11 rows with NaN in TotalCharges. Final dataset: 7,032 rows.',
    'Duplicate check: No duplicate rows found.',
    'No other missing values were detected across the remaining 20 columns.',
    'customerID column: Dropped as it is a unique identifier with no predictive value.',
    'Target variable: Encoded as binary (1 = Churned, 0 = Not Churned).',
    'Categorical variables: Label-encoded for model compatibility.',
]
for step in cleaning_steps:
    add_bullet(doc, step)
doc.add_paragraph()

# ── 7. Exploratory Data Analysis ─────────────────────────────────────────────
add_section_title(doc, '7. Exploratory Data Analysis', level=1)
add_body(doc,
    'EDA was conducted to understand churn patterns across key dimensions. '
    'Charts were created for churn distribution, contract type, internet service, '
    'payment method, tenure, and monthly charges.'
)
doc.add_paragraph()

add_section_title(doc, '7.1 Churn Distribution', level=2)
add_body(doc, 'Out of 7,032 customers: 5,163 did not churn (73.42%) and 1,869 churned (26.58%).')
doc.add_paragraph()

add_section_title(doc, '7.2 Churn by Contract Type', level=2)
add_table(doc,
    ['Contract Type', 'Churn Rate (%)'],
    [('Month-to-month', '42.71%'), ('One year', '11.28%'), ('Two year', '2.85%')]
)

add_section_title(doc, '7.3 Churn by Internet Service', level=2)
add_table(doc,
    ['Internet Service', 'Churn Rate (%)'],
    [('Fiber optic', '41.89%'), ('DSL', '19.00%'), ('No service', '7.43%')]
)

add_section_title(doc, '7.4 Churn by Payment Method', level=2)
add_table(doc,
    ['Payment Method', 'Churn Rate (%)'],
    [
        ('Electronic check', '45.29%'),
        ('Mailed check', '19.20%'),
        ('Bank transfer (automatic)', '16.73%'),
        ('Credit card (automatic)', '15.25%'),
    ]
)

add_section_title(doc, '7.5 Tenure and Monthly Charges', level=2)
add_table(doc,
    ['Metric', 'Churned', 'Not Churned'],
    [
        ('Average Tenure (months)', '17.98', '37.65'),
        ('Average Monthly Charges ($)', '74.44', '61.31'),
    ]
)

add_section_title(doc, '7.6 Churn Rate by Tenure Group', level=2)
add_body(doc,
    'Customers with tenure under 12 months have a churn rate of 47.68% — the highest of '
    'any tenure group. Churn rate declines consistently as tenure increases, dropping to '
    '6.61% for customers with 61–72 months of tenure.'
)
add_table(doc,
    ['Tenure Group (months)', 'Churn Rate (%)'],
    [
        ('0–12', '47.68%'),
        ('13–24', '28.71%'),
        ('25–36', '21.63%'),
        ('37–48', '19.03%'),
        ('49–60', '14.42%'),
        ('61–72', '6.61%'),
    ]
)

# ── 8. Machine Learning Models ────────────────────────────────────────────────
add_section_title(doc, '8. Machine Learning Models', level=1)

add_section_title(doc, '8.1 Feature Engineering', level=2)
eng_steps = [
    'Removed customerID (identifier column).',
    'Label-encoded all 15 categorical features using sklearn LabelEncoder.',
    'Applied StandardScaler to normalize features for Logistic Regression.',
    'Train-test split: 80% training (5,625 rows) / 20% testing (1,407 rows), stratified by Churn.',
]
for s in eng_steps:
    add_bullet(doc, s)
doc.add_paragraph()

add_section_title(doc, '8.2 Logistic Regression', level=2)
add_body(doc,
    'Logistic Regression is a linear classification algorithm that estimates the probability '
    'of a binary outcome. It was trained with max_iter=1000 and random_state=42 on scaled features.'
)
doc.add_paragraph()

add_section_title(doc, '8.3 Random Forest', level=2)
add_body(doc,
    'Random Forest is an ensemble of 100 decision trees (n_estimators=100, random_state=42). '
    'It operates on unscaled features and provides feature importance scores. '
    'The top predictors were: TotalCharges (18.49%), MonthlyCharges (17.72%), '
    'tenure (15.91%), Contract (8.30%), and PaymentMethod (5.01%).'
)
doc.add_paragraph()

# ── 9. Model Evaluation ───────────────────────────────────────────────────────
add_section_title(doc, '9. Model Evaluation', level=1)
add_body(doc,
    'Both models were evaluated on the held-out test set of 1,407 customers '
    '(374 churned, 1,033 not churned) using five metrics.'
)
doc.add_paragraph()

add_table(doc,
    ['Model', 'Accuracy', 'Precision', 'Recall', 'F1-Score', 'ROC-AUC'],
    [
        ('Logistic Regression', '79.39%', '62.43%', '56.42%', '59.27%', '83.45%'),
        ('Random Forest', '78.32%', '61.77%', '48.40%', '54.27%', '81.13%'),
    ]
)

add_section_title(doc, '9.1 Confusion Matrix — Logistic Regression', level=2)
add_table(doc,
    ['', 'Predicted: No Churn', 'Predicted: Churn'],
    [
        ('Actual: No Churn', '906', '127'),
        ('Actual: Churn', '163', '211'),
    ]
)

add_section_title(doc, '9.2 Confusion Matrix — Random Forest', level=2)
add_table(doc,
    ['', 'Predicted: No Churn', 'Predicted: Churn'],
    [
        ('Actual: No Churn', '921', '112'),
        ('Actual: Churn', '193', '181'),
    ]
)

# ── 10. Results ───────────────────────────────────────────────────────────────
add_section_title(doc, '10. Results', level=1)
add_body(doc,
    'Logistic Regression outperforms Random Forest across all key metrics, '
    'achieving a ROC-AUC of 83.45% and Recall of 56.42%. Recall is the most '
    'important metric for churn prediction — catching true churners early allows '
    'the business to intervene before customers leave. Logistic Regression correctly '
    'identified 211 churners out of 374 in the test set. The model is selected as '
    'the production model.'
)
doc.add_paragraph()

# ── 11. Business Recommendations ─────────────────────────────────────────────
add_section_title(doc, '11. Business Recommendations', level=1)
recs = [
    ('Promote Long-Term Contracts',
     'Month-to-month customers churn at 42.71% vs 2.85% for two-year customers. '
     'Offer discounts, loyalty rewards, or service upgrades to incentivize long-term commitments.'),
    ('Improve Fiber Optic Service',
     'Fiber optic subscribers churn at 41.89% — the highest rate. Investigate pricing strategy '
     'and service quality to address underlying satisfaction issues.'),
    ('Reduce Electronic Check Usage',
     'Electronic check customers churn at 45.29%. Offer incentives to switch to automatic payment '
     'methods (bank transfer or credit card), which correlate with significantly lower churn.'),
    ('Early Tenure Retention Programs',
     'Churned customers averaged only 17.98 months of tenure vs 37.65 months for retained customers. '
     'Implement structured onboarding, welcome programs, and proactive outreach in the first year.'),
    ('Deploy Churn Prediction System',
     'Use the Logistic Regression model (ROC-AUC: 83.45%) to score customers monthly. '
     'Prioritize retention outreach for customers with churn probability above 70%.'),
    ('Review High-Charge Customer Plans',
     'Churned customers paid $74.44/month on average vs $61.31 for retained customers. '
     'Consider personalized pricing or plan optimization for high-value customers showing churn signals.'),
]
for title, detail in recs:
    para = doc.add_paragraph(style='List Bullet')
    run_title = para.add_run(f'{title}: ')
    run_title.bold = True
    run_title.font.size = Pt(11)
    run_detail = para.add_run(detail)
    run_detail.font.size = Pt(11)
doc.add_paragraph()

# ── 12. Limitations ───────────────────────────────────────────────────────────
add_section_title(doc, '12. Limitations', level=1)
lims = [
    'The dataset is static (a single point-in-time snapshot). Real-world churn prediction benefits from time-series data.',
    'Label encoding assumes ordinal relationships between categories, which may not always be appropriate.',
    'Class imbalance (26.58% churn vs 73.42% not churned) was not explicitly addressed with techniques like SMOTE.',
    'The models were not tuned with hyperparameter optimization (e.g., GridSearchCV), which could improve performance.',
    'External factors (competitor pricing, market trends) are not captured in the dataset.',
]
for l in lims:
    add_bullet(doc, l)
doc.add_paragraph()

# ── 13. Future Scope ──────────────────────────────────────────────────────────
add_section_title(doc, '13. Future Scope', level=1)
future = [
    'Apply class-imbalance handling techniques (SMOTE, class_weight) to improve Recall.',
    'Perform hyperparameter tuning using GridSearchCV or RandomizedSearchCV.',
    'Explore advanced models: XGBoost, LightGBM, or neural networks.',
    'Build a time-series model using customer interaction history.',
    'Integrate the prediction model into a real-time CRM system for automated retention triggers.',
    'Use SHAP values for better model interpretability and individual customer explanations.',
]
for f in future:
    add_bullet(doc, f)
doc.add_paragraph()

# ── 14. Conclusion ────────────────────────────────────────────────────────────
add_section_title(doc, '14. Conclusion', level=1)
add_body(doc,
    'This project successfully built a customer churn prediction system for a telecom company '
    'using the IBM Telco Customer Churn dataset. After cleaning the data, performing exploratory '
    'analysis, and training two machine learning models, Logistic Regression was selected as '
    'the best model with an Accuracy of 79.39%, Recall of 56.42%, and ROC-AUC of 83.45%.'
)
doc.add_paragraph()
add_body(doc,
    'Key findings show that contract type, internet service type, and payment method are '
    'strongly associated with churn. Customers on month-to-month contracts, using fiber optic '
    'internet, and paying via electronic check are at the highest risk. Acting on these findings '
    'and deploying the predictive model can enable the company to proactively retain at-risk '
    'customers and improve long-term customer lifetime value.'
)
doc.add_paragraph()

# ── 15. References ────────────────────────────────────────────────────────────
add_section_title(doc, '15. References', level=1)
refs = [
    'IBM Sample Data Sets. Telco Customer Churn. Available at: https://www.ibm.com/communities/analytics/watson-analytics-blog/guide-to-sample-datasets/',
    'Kaggle. Telco Customer Churn Dataset. Available at: https://www.kaggle.com/datasets/blastchar/telco-customer-churn',
    'Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830.',
    'Breiman, L. (2001). Random Forests. Machine Learning, 45(1), 5-32.',
    'Cox, D.R. (1958). The Regression Analysis of Binary Sequences. Journal of the Royal Statistical Society, Series B, 20(2), 215-242.',
    'McKinney, W. (2010). Data Structures for Statistical Computing in Python. Proceedings of the 9th Python in Science Conference, 445, 51-56.',
    'Streamlit Inc. (2024). Streamlit Documentation. Available at: https://docs.streamlit.io/',
]
for i, ref in enumerate(refs, 1):
    para = doc.add_paragraph(style='List Number')
    para.add_run(ref).font.size = Pt(10)
doc.add_paragraph()

# Footer
doc.add_paragraph('─' * 80)
footer_para = doc.add_paragraph()
footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer_para.add_run(
    f'Telco Customer Churn Prediction | Machine Learning Project | {datetime.date.today().year}'
)
footer_run.font.size = Pt(9)
footer_run.font.color.rgb = RGBColor(0x57, 0x60, 0x6a)

# Save
doc.save('project_report.docx')
print('project_report.docx generated successfully.')
